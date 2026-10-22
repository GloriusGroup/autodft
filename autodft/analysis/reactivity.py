"""Reactivity singlepoints' properties as rows: molecular, atomic and Wiberg sheets.

One row set per relaxed or vertical singlepoint of a Reactivity state's
reported conformer (every valid conformer with *all_conformers*); excluded
conformers never appear. Atom indices are 1-based, as ORCA's NBO prints them.
"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from sqlmodel import Session, col, select

from autodft import categories
from autodft.db import get_session
from autodft.extraction import results
from autodft.extraction.extractor import ConformerResult, PipelineExtractor
from autodft.models import ComputationTask, Molecule, MoleculeState, TaskStatus, TaskType

SP_KINDS = {
    "singlepoint": "relaxed",
    "singlepoint_vert_ox": "vert_ox",
    "singlepoint_vert_red": "vert_red",
    "singlepoint_vert_spin_change": "vert_spin_change",
}
KEYS = ("mol_id", "smiles", "state", "state_id", "conformer", "opt_task_id", "sp_kind", "sp_task_id",
        "charge", "multiplicity")
QUADRUPOLE = (("xx", 0, 0), ("yy", 1, 1), ("zz", 2, 2), ("xy", 0, 1), ("xz", 0, 2), ("yz", 1, 2))
COLUMNS = {
    "molecular": (*KEYS, "energy_Eh", "polarizability_iso_au", *(f"quad_{q}_au" for q, _, _ in QUADRUPOLE),
                  "s2", "stable", "stability_lowest_Eh", "note"),
    "atomic": (*KEYS, "atom_1based", "element", "hirshfeld_charge", "hirshfeld_spin", "npa_charge", "npa_spin"),
    "wiberg": (*KEYS, "i_1based", "j_1based", "wiberg_index"),
}


def _flagged_states(session: Session, project_name: str) -> list[tuple[Molecule, MoleculeState]]:
    rows = session.exec(
        select(Molecule, MoleculeState)
        .join(MoleculeState, col(MoleculeState.molecule_id) == col(Molecule.id))
        .where(Molecule.project_name == project_name,
               col(MoleculeState.metadata_json).contains(categories.REACTIVITY))
        .order_by(col(Molecule.id), col(MoleculeState.id))
    ).all()
    return [(m, s) for m, s in rows if json.loads(s.metadata_json).get(categories.REACTIVITY)]


def has_reactivity(project_name: str) -> bool:
    """Whether any state of the project asked for Reactivity."""
    with get_session() as session:
        return bool(_flagged_states(session, project_name))


def reactivity_rows(project_name: str, all_conformers: bool = False) -> dict[str, list[dict]]:
    """``{molecular, atomic, wiberg}`` rows of every Reactivity state."""
    from autodft.engine.state_machine import _get_job_charge_multiplicity

    extractor = PipelineExtractor(project_name)
    # Planned under a short session, outputs parsed with none open (see db.py).
    plans = []
    with get_session() as session:
        for mol, state in _flagged_states(session, project_name):
            for conformer in _conformers(extractor, session, mol, state, all_conformers):
                tasks = session.exec(
                    select(ComputationTask).where(
                        ComputationTask.depends_on_task_id == conformer.opt_task_id,
                        ComputationTask.status == TaskStatus.successful,
                        col(ComputationTask.task_type).in_([TaskType(t) for t in SP_KINDS]),
                    )
                ).all()
                for task in sorted(tasks, key=lambda t: list(SP_KINDS).index(t.task_type.value)):
                    charge, multiplicity = _get_job_charge_multiplicity(task.task_type, state)
                    keys = {
                        "mol_id": mol.id, "smiles": mol.smiles, "state": state.description,
                        "state_id": state.id, "conformer": conformer.conformer_index,
                        "opt_task_id": conformer.opt_task_id, "sp_kind": SP_KINDS[task.task_type.value],
                        "sp_task_id": task.id, "charge": charge, "multiplicity": multiplicity,
                    }
                    job = results.latest_successful_job(session, task.id)
                    source = (job.job_path, *results.stored(session, job)) if job is not None else None
                    plans.append((keys, task.task_type.value, source, conformer.note))

    rows: dict[str, list[dict]] = {sheet: [] for sheet in COLUMNS}
    for keys, task_type, source, note in plans:
        record = None
        if source is not None:
            job_path, kept, current = source
            record = results.resolve(job_path, task_type, kept, current, require=("reactivity", "npa_charges"))
        _add(rows, keys, record, note)
    return rows


def _conformers(
    extractor: PipelineExtractor, session: Session, mol: Molecule, state: MoleculeState, all_conformers: bool,
) -> list[ConformerResult]:
    """The reported conformer, or every one not excluded."""
    found = extractor.extract_state_results(session, mol, state)
    if not all_conformers:
        picked = extractor._pick_reported_conformer(found)
        return [picked] if picked is not None else []
    return [c for c in found if not results.excludes(c.note)]


def _add(rows: dict[str, list[dict]], keys: dict, record: Optional[dict], note: Optional[str]) -> None:
    """One singlepoint's molecular row and its atomic and Wiberg rows."""
    reactivity = record.get("reactivity") if record else None
    notes = [n for n in (note, None if reactivity else "no reactivity output") if n]
    polar = (reactivity or {}).get("polarizability") or {}
    quad = ((reactivity or {}).get("quadrupole") or {}).get("tensor")
    stability = (reactivity or {}).get("stability") or {}
    rows["molecular"].append({
        **keys,
        "energy_Eh": record["energy"] if record else None,
        "polarizability_iso_au": polar.get("isotropic"),
        **{f"quad_{q}_au": quad[i][j] if quad else None for q, i, j in QUADRUPOLE},
        "s2": (reactivity or {}).get("s2"),
        "stable": stability.get("stable"),
        "stability_lowest_Eh": stability.get("lowest_eigenvalue"),
        "note": "; ".join(notes) or None,
    })
    if reactivity is None:
        return
    hirshfeld = {a["index"]: a for a in reactivity.get("hirshfeld") or []}
    npa = {a["index"]: a for a in record.get("npa_charges") or []}
    for index in sorted(hirshfeld.keys() | npa.keys()):
        h, n = hirshfeld.get(index, {}), npa.get(index, {})
        rows["atomic"].append({
            **keys, "atom_1based": index + 1, "element": h.get("element") or n.get("element"),
            "hirshfeld_charge": h.get("charge"), "hirshfeld_spin": h.get("spin"),
            "npa_charge": n.get("charge"), "npa_spin": n.get("spin"),
        })
    for i, j, value in reactivity.get("wiberg") or []:
        rows["wiberg"].append({**keys, "i_1based": i + 1, "j_1based": j + 1, "wiberg_index": value})


def build_xlsx(rows: dict[str, list[dict]]) -> bytes:
    """The three sheets, each with its header row even when empty."""
    from autodft.analysis.photophysics_export import _HEADER_FILL, _HEADER_FONT

    wb = Workbook()
    wb.remove(wb.active)
    for sheet, columns in COLUMNS.items():
        ws = wb.create_sheet(sheet)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
        for row in rows.get(sheet, []):
            ws.append([row.get(c) for c in columns])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
