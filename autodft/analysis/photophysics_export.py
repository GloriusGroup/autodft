"""The photophysics workbook: a sheet per category, plus the sticks and ESD jobs."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F3340")
RATES = ("isc", "risc", "ic", "fluorescence", "isc_t1_s0", "phosphorescence")
DERIVED = ("tau_s1_ns", "phi_fluorescence", "phi_isc", "phi_ic",
           "tau_t1_us", "phi_phosphorescence", "phi_isc_t1_s0", "phi_risc")


def build_xlsx(payload: dict) -> bytes:
    """Render ``spectroscopy.full_payload(...)`` as an XLSX workbook."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    for row in (
        ("Project", payload.get("project")),
        ("Molecules", len(payload.get("molecules", []))),
        ("Boltzmann temperature (K)", payload.get("temperature_k")),
        ("Generated (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
    ):
        summary.append(row)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 40

    uv, uv_sticks, ir, ir_sticks, nmr, esd, esd_jobs = [], [], [], [], [], [], []
    for m in payload.get("molecules", []):
        if m.get("uvvis"):
            e, peak = m["uvvis"], m["uvvis"].get("peak") or {}
            uv.append((m["id"], m["smiles"], m["state_id"], e["count"], e["pending"], e["failed"],
                       e["unavailable"], e["unweighted"], e["weighting"], peak.get("wavelength_nm"),
                       peak.get("energy_ev"), peak.get("fosc"), e.get("shortest_nm")))
            for c in e.get("conformers", []):
                for t in c["transitions"]:
                    uv_sticks.append((m["id"], m["state_id"], c["conformer_index"], c["weight"], t["root"],
                                      t["energy_ev"], t["wavelength_nm"], t["fosc"]))
        if m.get("ir"):
            e, peak = m["ir"], m["ir"].get("peak") or {}
            ir.append((m["id"], m["smiles"], m["state_id"], e["count"], e["pending"], e["failed"],
                       e["unavailable"], e["unweighted"], e["weighting"], peak.get("frequency_cm"),
                       peak.get("intensity_km_mol")))
            for c in e.get("conformers", []):
                for mode in c["modes"]:
                    ir_sticks.append((m["id"], m["state_id"], c["conformer_index"], c["weight"],
                                      mode["frequency_cm"], mode["intensity_km_mol"]))
        if m.get("nmr"):
            e = m["nmr"]
            for element, signals in (e.get("nuclei") or {}).items():
                ref = e.get("reference", {}).get(element, {})
                for s in signals:
                    nmr.append((m["id"], m["smiles"], m["state_id"], element, s["shift_ppm"], s["shielding_ppm"],
                                s["count"], ", ".join(str(a) for a in s["atoms"]), ref.get("compound"),
                                ref.get("status"), ref.get("method_matches"), e.get("equivalence")))
        if m.get("esd"):
            e, d = m["esd"], m["esd"].get("derived") or {}
            rates = e.get("rates", {})
            esd.append((m["id"], m["smiles"], m["state_id"], e["status"], e.get("temperature_k"),
                        e.get("herzberg_teller"), e.get("delta_est_ev"), e.get("delta_est_uks_ev"),
                        *(rates.get(r, {}).get("rate_s") for r in RATES),
                        *(d.get(k) for k in DERIVED), " | ".join(e.get("flags", []))))
            for r in RATES:
                for j in rates.get(r, {}).get("jobs", []):
                    esd_jobs.append((m["id"], m["state_id"], r, j.get("triplet"), j.get("sublevel"),
                                     j.get("rate_s"), j.get("dele_cm"), j.get("socme_cm"), j.get("fc_percent"),
                                     j.get("ht_percent"), j.get("k_squared"), j.get("e00_cm")))

    _sheet(wb, "UV-Vis", ("mol_id", "smiles", "state_id", "count", "pending", "failed", "unavailable",
                          "unweighted", "weighting", "peak_nm", "peak_eV", "peak_fosc", "shortest_nm"), uv)
    _sheet(wb, "UV-Vis sticks", ("mol_id", "state_id", "conformer", "weight", "root", "energy_eV",
                                 "wavelength_nm", "fosc"), uv_sticks)
    _sheet(wb, "IR", ("mol_id", "smiles", "state_id", "count", "pending", "failed", "unavailable",
                      "unweighted", "weighting", "peak_cm", "peak_intensity_km_mol"), ir)
    _sheet(wb, "IR sticks", ("mol_id", "state_id", "conformer", "weight", "frequency_cm",
                             "intensity_km_mol"), ir_sticks)
    _sheet(wb, "NMR", ("mol_id", "smiles", "state_id", "nucleus", "shift_ppm", "shielding_ppm", "count",
                       "atoms", "reference", "reference_status", "method_matches", "equivalence"), nmr)
    _sheet(wb, "ESD", ("mol_id", "smiles", "state_id", "status", "temperature_K", "HT", "dEST_eV",
                       "dEST_UKS_eV", "k_ISC", "k_RISC", "k_IC", "k_F", "k_ISC_T1S0", "k_P", "tau_S1_ns",
                       "phi_F", "phi_ISC", "phi_IC", "tau_T1_us", "phi_P", "phi_ISC_T1S0", "phi_RISC",
                       "flags"), esd)
    _sheet(wb, "ESD jobs", ("mol_id", "state_id", "rate", "triplet", "sublevel", "rate_s", "dele_cm",
                            "socme_cm", "fc_percent", "ht_percent", "k_squared", "e00_cm"), esd_jobs)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sheet(wb: Workbook, title: str, header: tuple, rows: list) -> None:
    """A sheet with a styled header row; none when there is nothing to show."""
    if not rows:
        return
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append(list(row))
