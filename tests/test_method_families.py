"""Method families (a molecule's states at one header set) kept apart in state analysis and exports."""

from __future__ import annotations

import csv
import io
import json
import shutil
import subprocess

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from autodft.analysis import state_analysis
from autodft.analysis.state_analysis import analyze_project
from autodft.db import get_session
from autodft.extraction.extractor import PipelineExtractor
from autodft.models import MoleculeState, TaskType
from tests.test_dashboard import _functions
from tests.test_export_tree import (  # noqa: F401 - db is a fixture
    PROJECT, _conformer, _export, _molecule, _root, _standard, _task, _tree, db,
)

FAMILY = ["S0_state_id", "confsearch_header_id", "optimization_header_id", "singlepoint_header_id"]


def _state(session, mol_id, description, headers, multiplicity=1, charge=0) -> MoleculeState:
    state = MoleculeState(molecule_id=mol_id, description=description, multiplicity=multiplicity, charge=charge,
                          metadata_json=json.dumps({"imag_freq_tolerance": -10.0}),
                          confsearch_header_id=headers[0], optimization_header_id=headers[1],
                          singlepoint_header_id=headers[2])
    session.add(state)
    session.commit()
    return state


def _ox(session, settings, state, e_sp) -> None:
    opt = _task(session, settings, state, TaskType.optimization, f"ox_opt_{state.id}",
                {"energy": None, "free_energy_correction": 0.1, "imaginary": [], "followed_root": None})
    _task(session, settings, state, TaskType.singlepoint, f"ox_sp_{state.id}", {"energy": e_sp}, depends_on=opt.id)


def _two_families(settings, same_headers=False) -> dict:
    """S0 at headers A (G -99.9) with ox (G -99.65), and S0 at headers B (G -100.9) without ox."""
    a_headers, b_headers = (None, 1, 2), (None, 1, 2) if same_headers else (None, 3, 4)
    with get_session() as session:
        a = _molecule(session, {"imag_freq_tolerance": -10.0})
        a.optimization_header_id, a.singlepoint_header_id = a_headers[1:]
        session.add(a)
        session.commit()
        _task(session, settings, a, TaskType.confsearch, "cs_a")
        _conformer(session, settings, a, "a", -100.0, uvvis=False)
        b = _state(session, a.molecule_id, "S0", b_headers)
        _task(session, settings, b, TaskType.confsearch, "cs_b")
        _conformer(session, settings, b, "b", -101.0, uvvis=False)
        ox = _state(session, a.molecule_id, "ox", a_headers, multiplicity=2, charge=1)
        _ox(session, settings, ox, -99.75)
        return {"mol": a.molecule_id, "a": a.id, "b": b.id, "ox": ox.id}


def test_each_family_gets_its_own_analysis(db):
    ids = _two_families(db)
    payload = analyze_project(PROJECT, use_cache=False)
    rows = payload["molecules"]
    assert [(m["id"], m["state_id"], m["header_ids"]) for m in rows] == [
        (ids["mol"], ids["a"], [None, 1, 2]), (ids["mol"], ids["b"], [None, 3, 4])]
    for mode in ("lowest_energy", "rmsd_matched"):
        a, b = (m["modes"][mode] for m in rows)
        assert a["ox"]["dG_h"] == pytest.approx(0.25)
        assert set(a["picks"]) == {"S0", "ox"}
        assert b["ox"]["dG_h"] is None and set(b["picks"]) == {"S0"}
        assert b["picks"]["S0"]["e_combined"] == pytest.approx(-100.9)


def test_the_workbook_names_each_family(db):
    ids = _two_families(db)
    wb = load_workbook(io.BytesIO(state_analysis.build_xlsx_bytes(analyze_project(PROJECT, use_cache=False))))
    sheet = [list(r) for r in wb["Lowest Energy"].iter_rows(values_only=True)]
    assert sheet[0][-4:] == FAMILY
    assert [r[-4:] for r in sheet[1:]] == [[ids["a"], None, 1, 2], [ids["b"], None, 3, 4]]
    ox_dg = sheet[0].index("ox_dG_Eh")
    assert sheet[1][ox_dg] == pytest.approx(0.25) and sheet[2][ox_dg] is None
    conformers = [list(r) for r in wb["Conformers"].iter_rows(values_only=True)]
    assert conformers[0][-4:] == FAMILY
    assert {(r[3], r[-4]) for r in conformers[1:]} == {("S0", ids["a"]), ("ox", ids["a"]), ("S0", ids["b"])}
    assert wb["Summary"]["B7"].value == 1


def test_the_csv_names_each_rows_state_and_family(db):
    ids = _two_families(db)
    _export(db)
    with open(_root(db) / "p.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert [(r["state"], r["state_id"], r["confsearch_header_id"], r["optimization_header_id"],
             r["singlepoint_header_id"], r["e_combined"]) for r in rows] == [
        ("S0", str(ids["a"]), "", "1", "2", "-99.9"),
        ("S0", str(ids["b"]), "", "3", "4", "-100.9"),
        ("ox", str(ids["ox"]), "", "1", "2", "-99.65"),
    ]


def test_the_json_summary_names_each_rows_state(db, tmp_path):
    ids = _two_families(db)
    PipelineExtractor(PROJECT).export_summary_json(tmp_path / "s.json")
    rows = json.loads((tmp_path / "s.json").read_text())
    assert [(r["state"], r["state_id"], r["optimization_header_id"]) for r in rows] == [
        ("S0", ids["a"], 1), ("S0", ids["b"], 3), ("ox", ids["ox"], 1)]


def test_a_state_unknown_to_the_summary_leaves_its_headers_empty(db, tmp_path, monkeypatch):
    """A state that arrives while the summary is written has no headers there, and no KeyError."""
    ids = _two_families(db)
    original = PipelineExtractor.extract_results

    def late(self, all_conformers=False):
        rows = original(self, all_conformers)
        rows[-1].state_id = 9999
        return rows
    monkeypatch.setattr(PipelineExtractor, "extract_results", late)
    PipelineExtractor(PROJECT).export_summary_json(tmp_path / "s.json")
    rows = json.loads((tmp_path / "s.json").read_text())
    assert [(r["state_id"], r["optimization_header_id"]) for r in rows] == [(ids["a"], 1), (ids["b"], 3), (9999, None)]


@pytest.mark.parametrize("same_headers", [False, True])
def test_states_sharing_a_description_get_their_own_folders(db, same_headers):
    ids = _two_families(db, same_headers)
    _export(db)
    a, b = f"S0__state{ids['a']}", f"S0__state{ids['b']}"
    tree = _tree(db, ids["mol"])
    assert {p.split("/")[0] for p in tree if "/" in p} == {a, b, "ox"}
    folder = _root(db) / "molecules" / str(ids["mol"])
    assert (folder / a / "conf1" / "opt.out").read_text() == "opt_a job_1 output.out"
    assert (folder / b / "conf1" / "opt.out").read_text() == "opt_b job_1 output.out"
    assert (folder / a / "confsearch.out").read_text() == "cs_a job_1 output.out"
    assert (folder / b / "confsearch.out").read_text() == "cs_b job_1 output.out"
    manifest = json.loads((folder / "molecule.json").read_text())
    assert [[c["folder"] for c in s["conformers"]] for s in manifest["states"]] == [
        [f"{a}/conf1"], [f"{b}/conf1"], ["ox/conf1"]]


def test_one_header_set_is_one_family(db):
    """The NMR reference project's S0s per NMR method share headers: one analysis, separate folders."""
    ids = _two_families(db, same_headers=True)
    [row] = analyze_project(PROJECT, use_cache=False)["molecules"]
    assert row["state_id"] == ids["a"]
    assert row["modes"]["lowest_energy"]["picks"]["S0"]["e_combined"] == pytest.approx(-100.9)


def test_an_archived_project_keeps_its_families_apart(db):
    ids = _two_families(db)
    PipelineExtractor(PROJECT).archive_project(db.export_data_path, db.comp_data_path, False, False)
    payload = analyze_project(PROJECT, use_cache=False)
    assert payload["archive_source"].endswith("p_all_conformers.csv")
    a, b = payload["molecules"]
    assert (a["state_id"], b["state_id"]) == (ids["a"], ids["b"])
    assert a["modes"]["lowest_energy"]["ox"]["dG_h"] == pytest.approx(0.25)
    assert b["modes"]["lowest_energy"]["ox"]["dG_h"] is None


# Written by the code before method families, from _standard(ox=True) with every conformer.
_CSV_HEAD = ("molecule_id,smiles,state,conformer_index,opt_task_id,e_singlepoint,e_correction,e_combined,"
             "e_vert_spin_change,e_vert_ox,e_vert_red,note\n")
BEFORE = {
    "p.csv": _CSV_HEAD + "1,c1ccccc1,S0,1,2,-2.0,0.1,-1.9,,-1.8,,\n1,c1ccccc1,ox,1,10,-1.5,0.1,-1.4,,,,\n",
    "p_all_conformers.csv": _CSV_HEAD + "1,c1ccccc1,S0,1,2,-2.0,0.1,-1.9,,-1.8,,\n"
                            "1,c1ccccc1,S0,2,6,-1.0,0.1,-0.9,,-0.8,,\n1,c1ccccc1,ox,1,10,-1.5,0.1,-1.4,,,,\n",
}
BEFORE_TREE = {"molecule.json", "S0/confsearch.out", "ox/conf1/opt.inp", "ox/conf1/opt.out", "ox/conf1/opt.xyz",
               "ox/conf1/sp.inp", "ox/conf1/sp.out"} | {
    f"S0/conf{k}/{name}" for k in (1, 2) for name in (
        "opt.inp", "opt.out", "opt.xyz", "sp.inp", "sp.out", "uvvis.inp", "uvvis.out", "vert_ox.inp", "vert_ox.out")}
BEFORE_MODE_SHEET = [
    ["mol_id", "smiles", "S0_conf", "T1_conf", "T1_rmsd_to_S0_A", "ox_conf", "ox_rmsd_to_S0_A", "red_conf",
     "red_rmsd_to_S0_A", "triplet_dE_adiab_Eh", "triplet_lambda_fwd_Eh", "triplet_lambda_bwd_Eh", "ox_dG_Eh",
     "ox_E_vs_SCE_V", "ox_lambda_fwd_Eh", "ox_lambda_bwd_Eh", "red_dG_Eh", "red_E_vs_SCE_V", "red_lambda_fwd_Eh",
     "red_lambda_bwd_Eh"],
    [1, "c1ccccc1", 1, None, None, 1, None, None, None, None, None, None, 0.5, None, -0.3,
     None, None, None, None, None],
]
BEFORE_CONFORMERS_HEAD = ["mode", "mol_id", "smiles", "state", "conformer_index", "opt_task_id", "rmsd_to_S0_A",
                          "e_sp_Eh", "e_correction_Eh", "e_combined_Eh", "e_vert_ox_Eh", "e_vert_red_Eh",
                          "e_vert_spin_change_Eh"]


def test_a_single_family_project_exports_as_before(db):
    mol_id, _, _ = _standard(db, ox=True)
    with get_session() as session:
        for state in session.exec(select(MoleculeState)).all():
            state.confsearch_header_id, state.optimization_header_id, state.singlepoint_header_id = 2, 3, 4
            session.add(state)
        session.commit()
    _export(db, all_conformers=True)
    for name, text in BEFORE.items():
        assert (_root(db) / name).read_text() == text
    assert _tree(db, mol_id) == BEFORE_TREE
    manifest = json.loads((_root(db) / "molecules" / str(mol_id) / "molecule.json").read_text())
    assert [[c["folder"] for c in s["conformers"]] for s in manifest["states"]] == [["S0/conf1", "S0/conf2"],
                                                                                   ["ox/conf1"]]
    wb = load_workbook(_root(db) / "p_state_analysis.xlsx")
    for title in ("Lowest Energy", "RMSD Matched"):
        assert [list(r) for r in wb[title].iter_rows(values_only=True)] == BEFORE_MODE_SHEET
    assert [c.value for c in wb["Conformers"][1]] == BEFORE_CONFORMERS_HEAD
    assert wb["Summary"]["B7"].value == 1


def _cards(payload: dict) -> list[str]:
    script = ("var SA_UNITS = {kcal: {label: 'kcal/mol', scale: 627.5, digits: 2}};\n"
              "var host = {innerHTML: ''}; var document = {getElementById: function () { return host; }};\n"
              + _functions("escHtml", "fmtFixed", "fmtEnergy", "fmtVolts", "fmtRmsd", "renderSaColumn",
                           "saFamilyLabel", "renderSaCards")
              + f"\nrenderSaCards({json.dumps(payload)}, 'lowest_energy', 'kcal');\n"
              + "process.stdout.write(host.innerHTML);\n")
    out = subprocess.run(["node", "-e", script], capture_output=True, text=True, check=True).stdout
    return out.split('<div class="sa-card">')[1:]


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
def test_the_dashboard_labels_a_molecules_families(db):
    ids = _two_families(db)
    first, second = _cards(analyze_project(PROJECT, use_cache=False))
    assert f"S0 state {ids['a']} · headers –/1/2" in first
    assert f"S0 state {ids['b']} · headers –/3/4" in second


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
def test_the_dashboard_card_of_a_single_family_is_unlabelled(db):
    _standard(db, ox=True)
    [card] = _cards(analyze_project(PROJECT, use_cache=False))
    assert "S0 state" not in card
