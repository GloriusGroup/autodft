"""E vs SCE (MeCN) only when the opt and singlepoint of both states in ΔG ran in MeCN."""

from __future__ import annotations

import io
import json
import shutil

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from autodft.analysis.state_analysis import HARTREE_TO_EV, SCE_ABS_MECN, analyze_project, build_xlsx_bytes
from autodft.db import get_session
from autodft.models import ComputationHeader, MoleculeState
from tests.test_export_tree import PROJECT, _conformer, _molecule, db  # noqa: F401 - db is a fixture
from tests.test_method_families import _cards, _ox, _state

E_OX = 0.25 * HARTREE_TO_EV - SCE_ABS_MECN
MECN, GAS = ("CPCM", "MeCN"), ("omit", "")


def _metadata(opt, sp) -> dict:
    stage = lambda s: {"solvation": s[0], "solvent": s[1], "maxcore": 1000, "nprocs": 4}  # noqa: E731
    return {"imag_freq_tolerance": -10.0, "advanced": {"optimization": stage(opt), "singlepoint": stage(sp)}}


def _redox(settings, s0_solvation, ox_solvation=None) -> int:
    """S0 (G -99.9) and ox (G -99.65) at one header set, each with its own (opt, SP) solvation; the opt header id."""
    with get_session() as session:
        headers = [ComputationHeader(header_text=f"!B3LYP {basis}\n") for basis in ("def2-SVP", "def2-TZVP")]
        session.add_all(headers)
        session.commit()
        opt_id, sp_id = (h.id for h in headers)
        s0 = _molecule(session, _metadata(*s0_solvation))
        s0.optimization_header_id, s0.singlepoint_header_id = opt_id, sp_id
        session.add(s0)
        session.commit()
        _conformer(session, settings, s0, "a", -100.0, uvvis=False)
        ox = _state(session, s0.molecule_id, "ox", (None, opt_id, sp_id), multiplicity=2, charge=1)
        ox.metadata_json = json.dumps(_metadata(*(ox_solvation or s0_solvation)))
        session.add(ox)
        session.commit()
        _ox(session, settings, ox, -99.75)
        return opt_id


def _ox_result(mode="lowest_energy") -> dict:
    [row] = analyze_project(PROJECT, use_cache=False)["molecules"]
    return row["modes"][mode]["ox"]


NOTE = "E vs SCE needs opt and singlepoint in MeCN; "


@pytest.mark.parametrize("opt, sp, potential", [
    (MECN, GAS, None),
    (MECN, MECN, pytest.approx(E_OX)),
    (GAS, MECN, None),
    (GAS, GAS, None),
])
def test_only_opt_and_singlepoint_both_in_mecn_give_a_potential(db, opt, sp, potential):
    _redox(db, (opt, sp))
    for mode in ("lowest_energy", "rmsd_matched"):
        ox = _ox_result(mode)
        assert ox["dG_h"] == pytest.approx(0.25)
        assert ox["E_vs_sce_v"] == potential
    assert analyze_project(PROJECT, use_cache=False)["solvation_mecn"] is (opt == sp == MECN)


@pytest.mark.parametrize("opt, sp, note", [
    (MECN, GAS, "S0 singlepoint and ox singlepoint are not in MeCN"),
    (GAS, MECN, "S0 opt and ox opt are not in MeCN"),
])
def test_mixed_solvation_says_why(db, opt, sp, note):
    _redox(db, (opt, sp))
    for mode in ("lowest_energy", "rmsd_matched"):
        assert _ox_result(mode)["note"] == NOTE + note


@pytest.mark.parametrize("opt, sp", [(MECN, MECN), (GAS, GAS)])
def test_no_note_without_mixed_solvation(db, opt, sp):
    _redox(db, (opt, sp))
    assert "note" not in _ox_result()


@pytest.mark.parametrize("ox_solvation, note", [
    ((MECN, GAS), "ox singlepoint is not in MeCN"),
    ((GAS, MECN), "ox opt is not in MeCN"),
])
def test_both_picks_must_be_in_mecn(db, ox_solvation, note):
    """S0 fully in MeCN, ox with other Advanced settings: no potential."""
    _redox(db, (MECN, MECN), ox_solvation=ox_solvation)
    ox = _ox_result()
    assert ox["dG_h"] == pytest.approx(0.25) and ox["E_vs_sce_v"] is None
    assert ox["note"] == NOTE + note


def test_a_gas_molecule_in_a_mecn_project_says_why(db):
    _redox(db, (MECN, MECN))
    _redox(db, (GAS, GAS))
    mecn, gas = (m["modes"]["lowest_energy"]["ox"] for m in analyze_project(PROJECT, use_cache=False)["molecules"])
    assert mecn["E_vs_sce_v"] == pytest.approx(E_OX) and "note" not in mecn
    assert gas["E_vs_sce_v"] is None and gas["note"] == NOTE + "neither S0 nor ox ran in MeCN"


def test_a_legacy_header_is_read_as_written(db):
    """Without Advanced settings the header text itself decides."""
    opt_id = _redox(db, (GAS, GAS))
    with get_session() as session:
        session.get(ComputationHeader, opt_id).header_text = "!B3LYP def2-SVP CPCM(MeCN)\n"
        for state in session.exec(select(MoleculeState)).all():
            state.metadata_json = json.dumps({"imag_freq_tolerance": -10.0})
        session.commit()
    ox = _ox_result()
    assert ox["E_vs_sce_v"] is None and ox["note"] == NOTE + "S0 singlepoint and ox singlepoint are not in MeCN"


def _sheet(title="Lowest Energy") -> list[list]:
    wb = load_workbook(io.BytesIO(build_xlsx_bytes(analyze_project(PROJECT, use_cache=False))))
    return [list(r) for r in wb[title].iter_rows(values_only=True)]


def test_the_workbook_gives_the_reason(db):
    _redox(db, (GAS, MECN))
    for title in ("Lowest Energy", "RMSD Matched"):
        head, row = _sheet(title)
        assert head[-2:] == ["ox_sce_note", "red_sce_note"]
        assert row[-2:] == [NOTE + "S0 opt and ox opt are not in MeCN", None]
        assert row[head.index("ox_E_vs_SCE_V")] is None


def test_a_workbook_without_reasons_keeps_its_columns(db):
    _redox(db, (MECN, MECN))
    head, row = _sheet()
    assert head[-1] == "red_lambda_bwd_Eh" and "ox_sce_note" not in head
    assert row[head.index("ox_E_vs_SCE_V")] == pytest.approx(E_OX)


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
def test_the_dashboard_shows_the_note(db):
    _redox(db, (MECN, GAS))
    [card] = _cards(analyze_project(PROJECT, use_cache=False))
    assert NOTE + "S0 singlepoint and ox singlepoint are not in MeCN" in card
