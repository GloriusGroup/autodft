"""ESD results read back per molecule (real glyoxal outputs)."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from sqlmodel import select

from autodft import categories
from autodft.analysis.esd import EH_TO_EV, molecule_esd
from autodft.config import Settings
from autodft.engine import photophysics
from autodft.engine.state_machine import _create_job_for_task
from autodft.extraction import results
from autodft.extraction.extractor import PipelineExtractor
from autodft.models import ComputationJob, ComputationTask, TaskStatus, TaskType
from autodft.qm.orca import esd_inputs
from autodft.qm.orca.parser import OrcaParser
from tests.test_esd_joins import FIXTURES, _job, _rates, esd  # noqa: F401 - fixture
from tests.test_esd_parser import DONE

OUTPUTS = {
    TaskType.esd_isc: "isc_two_channels.out", TaskType.esd_risc: "risc_ht.out",
    TaskType.esd_ic: "ic.out", TaskType.esd_fluor: "fluor.out",
    TaskType.esd_isc_t1s0: "t1s0.out", TaskType.esd_phosp: "phosp.out",
}
E_UKS_T1 = -227.466493


def _uks_singlepoint(session, esd):
    opt = esd["tasks"]["T1"][0]
    sp = ComputationTask(task_type=TaskType.singlepoint, status=TaskStatus.successful,
                         state_id=esd["states"]["T1"].id, header_id=opt.header_id,
                         depends_on_task_id=opt.id, has_followups=False)
    session.add(sp)
    session.commit()
    _job(session, sp, esd["tmp_path"] / "jobs" / "sp_T1", f"FINAL SINGLE POINT ENERGY      {E_UKS_T1}\n")


def _finish_rates(session, esd):
    # Two ISC channels (T1, T2), matching isc_two_channels.out.
    s1 = esd["states"]["S1"]
    s1.metadata_json = json.dumps({**json.loads(s1.metadata_json), "esd_tn_window_ev": 1.0})
    session.add(s1)
    session.commit()
    photophysics.advance_photophysics(session, Settings())
    for task in _rates(session, esd).values():
        task.task_path = str(esd["tmp_path"] / "rates" / task.task_type.value)
        job = _create_job_for_task(session, task, 1, Settings(), qm_engine=OrcaParser())
        (Path(job.job_path) / "output.out").write_text(
            _with_diagnostics(task, (FIXTURES / OUTPUTS[task.task_type]).read_text()))
        job.success = True
        task.status = TaskStatus.successful
        session.add_all([job, task])
    session.commit()


def _diagnostic_chunk(rate: float) -> str:
    return f"The calculated ISC rate constant is\t\t{rate:e} s-1\n{DONE}"


def _with_diagnostics(task, text: str, rates: dict = None) -> str:
    """*text* (the main jobs' output) plus one FC rate per separate diagnostic job."""
    computed = json.loads(task.inputs_json)["computed"]
    extra = sorted({d["job"]: d for d in computed.get("diagnostics", [])
                    if d["job"] > len(computed["jobs"])}.values(), key=lambda d: d["job"])
    rates = rates or {}
    return text + "".join(_diagnostic_chunk(rates.get((d["triplet"], d["variant"]), 1.0e4)) for d in extra)


@pytest.fixture()
def done(session, esd):
    _uks_singlepoint(session, esd)
    _finish_rates(session, esd)
    return esd


def _analyse(session, esd, detail=True):
    return molecule_esd(session, PipelineExtractor("nho/p"), esd["states"]["S0"], detail)


def test_rates(session, done):
    result = _analyse(session, done)
    rates = result["rates"]
    assert result["status"] == "done"
    assert rates["isc"]["rate_s"] == pytest.approx(9.521341e3)  # T2's -1.5e-9 is set to 0
    assert rates["risc"]["rate_s"] == pytest.approx(3.577340e-05)
    assert rates["ic"]["rate_s"] == pytest.approx(2.646887e4)
    assert rates["fluorescence"]["rate_s"] == pytest.approx(4.811336e3)
    assert rates["isc_t1_s0"]["rate_s"] == pytest.approx(0.1598812)
    assert rates["phosphorescence"]["rate_s"] == pytest.approx((1.187214 + 9.715617e-02 + 1.328943e2) / 3)
    assert rates["fluorescence"]["e00_ev"] == pytest.approx(19321.27 / 8065.544)
    assert any("negative rate" in f for f in result["flags"])


def test_isc_jobs_carry_the_inputs(session, done):
    jobs = _analyse(session, done)["rates"]["isc"]["jobs"]
    assert [j["triplet"] for j in jobs] == [1, 2]
    assert jobs[0]["dele_cm"] == pytest.approx(5369.5, abs=0.05) and jobs[0]["socme_cm"] == pytest.approx(0.88)
    assert jobs[1]["rate_s"] == 0.0


def test_energy_gaps_and_socmes(session, done):
    result = _analyse(session, done)
    assert result["delta_est_ev"] == pytest.approx(0.665736, abs=1e-5)
    assert result["delta_est_uks_ev"] == pytest.approx(0.507794, abs=1e-5)
    assert result["socme_cm"] == {"S1_T1_at_T1": pytest.approx(0.88), "S1_T1_at_S1": pytest.approx(0.89),
                                  "T1_S0_at_S0": pytest.approx(0.02)}


def test_lifetimes_and_yields(session, done):
    derived = _analyse(session, done)["derived"]
    k_s1 = 9.521341e3 + 4.811336e3 + 2.646887e4
    assert derived["tau_s1_ns"] == pytest.approx(1e9 / k_s1)
    assert derived["phi_fluorescence"] == pytest.approx(4.811336e3 / k_s1)
    assert derived["phi_isc"] + derived["phi_fluorescence"] + derived["phi_ic"] == pytest.approx(1.0)
    k_p = (1.187214 + 9.715617e-02 + 1.328943e2) / 3
    k_t1 = k_p + 0.1598812 + 3.577340e-05
    assert derived["tau_t1_us"] == pytest.approx(1e6 / k_t1)
    assert derived["phi_phosphorescence"] == pytest.approx(k_p / k_t1)


@pytest.fixture()
def done_without_ic(session, esd):
    s1 = esd["states"]["S1"]
    s1.metadata_json = json.dumps({**json.loads(s1.metadata_json), "esd_ic": False})
    session.add(s1)
    session.commit()
    _uks_singlepoint(session, esd)
    _finish_rates(session, esd)
    return esd


def test_ic_not_requested(session, done_without_ic):
    result = _analyse(session, done_without_ic)
    assert result["status"] == "done"
    assert result["rates"]["ic"] == {"status": "not requested"}
    derived = result["derived"]
    k_s1 = 9.521341e3 + 4.811336e3
    assert derived["tau_s1_ns"] == pytest.approx(1e9 / k_s1)
    assert derived["phi_fluorescence"] + derived["phi_isc"] == pytest.approx(1.0)
    assert "phi_ic" not in derived
    assert "τ(S1) and the S1 yields exclude internal conversion (not requested)." in result["flags"]


def test_an_existing_ic_task_is_shown_without_the_option(session, done):
    s1 = done["states"]["S1"]
    s1.metadata_json = json.dumps({k: v for k, v in json.loads(s1.metadata_json).items() if k != "esd_ic"})
    session.add(s1)
    session.commit()
    result = _analyse(session, done)
    assert result["rates"]["ic"]["rate_s"] == pytest.approx(2.646887e4)
    assert "phi_ic" in result["derived"]
    assert not any("internal conversion" in f for f in result["flags"])


def test_a_molecule_from_before_the_option_still_expects_ic(session, esd):
    s1 = esd["states"]["S1"]
    s1.metadata_json = json.dumps({k: v for k, v in json.loads(s1.metadata_json).items() if k != "esd_ic"})
    session.add(s1)
    session.commit()
    assert _analyse(session, esd)["rates"]["ic"]["status"] != "not requested"


def test_the_workbook_says_not_requested(session, done_without_ic):
    from io import BytesIO

    from openpyxl import load_workbook

    from autodft.analysis.photophysics_export import build_xlsx

    esd = _analyse(session, done_without_ic)
    payload = {"project": "nho/p", "molecules": [{"id": 1, "smiles": "O=CC=O", "state_id": 1, "esd": esd}]}
    rows = list(load_workbook(BytesIO(build_xlsx(payload)))["ESD"].iter_rows(values_only=True))
    row = dict(zip(rows[0], rows[1]))
    assert row["k_IC"] == "not requested" and row["phi_IC"] is None
    assert row["k_F"] == pytest.approx(4.811336e3)


def test_the_summary_leaves_out_the_details(session, done):
    result = _analyse(session, done, detail=False)
    assert "jobs" not in result["rates"]["isc"] and "socme_cm" not in result


def test_a_failed_s1_optimisation_blocks_its_rates(session, esd):
    opt = esd["tasks"]["S1"][0]
    opt.status = TaskStatus.failed
    session.add(opt)
    session.add(ComputationJob(task_id=opt.id, attempt=2, success=False,
                               fail_reason="['Termination', 'LibXC Needed']"))
    session.commit()
    photophysics.advance_photophysics(session, Settings())
    rates = _analyse(session, esd)["rates"]
    assert rates["isc"]["status"] == "blocked"
    assert "S1 optimisation" in rates["isc"]["reason"] and "LibXC" in rates["isc"]["reason"]
    assert rates["phosphorescence"]["status"] == "created"


def test_a_dissociating_s1_is_explained(session, esd):
    opt = esd["tasks"]["S1"][0]
    opt.status = TaskStatus.failed
    session.add(opt)
    session.add(ComputationJob(task_id=opt.id, attempt=1, success=False,
                               fail_reason="['Termination', 'S1 Stays Bound']"))
    session.commit()
    photophysics.advance_photophysics(session, Settings())
    assert "S1 dissociates" in _analyse(session, esd)["rates"]["isc"]["reason"]


def test_a_recovered_input_stops_reporting_waiting(session, esd):
    """I1: once esd_done no longer latches on a partial join, a rate whose
    input recovers is created instead of reporting 'waiting' forever."""
    soc = esd["tasks"]["S0"][1]
    soc.status = TaskStatus.failed
    session.add(soc)
    session.commit()
    photophysics.advance_photophysics(session, Settings())
    assert _analyse(session, esd)["rates"]["ic"]["status"] == "blocked"

    soc.status = TaskStatus.successful
    session.add(soc)
    session.commit()
    photophysics.advance_photophysics(session, Settings())
    rates = _analyse(session, esd)["rates"]
    for name in ("ic", "fluorescence", "isc_t1_s0", "phosphorescence"):
        assert rates[name]["status"] != "waiting"


def test_before_seeding_and_after_a_seed_error(session, esd):
    s1 = esd["states"]["S1"]
    s1.metadata_json = json.dumps({"esd_role": "S1"})
    session.add(s1)
    session.commit()
    assert _analyse(session, esd)["status"] == "waiting"
    s1.metadata_json = json.dumps({"esd_role": "S1", "esd_seed": {"error": "no conformer"}})
    session.add(s1)
    session.commit()
    result = _analyse(session, esd)
    assert (result["status"], result["reason"]) == ("failed", "no conformer")


def test_risc_far_uphill_is_flagged(session, done):
    # glyoxal delta_est_ev ~= 0.6657 eV ~= 26 kT at 298.15 K.
    result = _analyse(session, done)
    assert any("risc" in f and "kT above T1" in f for f in result["flags"])


def test_risc_close_to_thermal_is_not_flagged(session, done, monkeypatch):
    fake = {"S0": 0.0, "S1": 0.1 / EH_TO_EV, "T1": 0.0}
    monkeypatch.setattr(esd_inputs, "energy", lambda state, data: fake[state])
    result = _analyse(session, done)
    assert result["delta_est_ev"] == pytest.approx(0.1)
    assert not any("kT above T1" in f for f in result["flags"])


def test_a_drifted_s1_root_is_flagged_in_detail(session, done):
    path = done["tmp_path"] / "jobs" / "opt_S1" / "output.out"
    path.write_text("DE(CIS) =      0.087063157 Eh (Root  2)\n")
    result = _analyse(session, done, detail=True)
    assert any("followed root 2" in f for f in result["flags"])
    assert not any("followed root" in f for f in _analyse(session, done, detail=False)["flags"])


def test_a_large_displacement_is_flagged(session, done):
    task = _rates(session, done)[TaskType.esd_fluor]
    path = Path(session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one().job_path)
    text = (path / "output.out").read_text().replace("0.739141", "8.500000")
    (path / "output.out").write_text(text)
    assert any("K*K" in f for f in _analyse(session, done)["flags"])


def _store_every_successful_job(session):
    """Backfill, scoped to the fixture's own engine (not autodft.db.get_session())."""
    jobs = session.exec(select(ComputationJob).where(ComputationJob.success == True)).all()  # noqa: E712
    for job in jobs:
        task = session.get(ComputationTask, job.task_id)
        record = results.parse(session, task, Path(job.job_path))
        if record is not None:
            results.save(session, task, job, record)
    session.commit()


def test_stored_records_give_the_same_payload(session, done):
    before = _analyse(session, done)
    _store_every_successful_job(session)
    assert _analyse(session, done) == before

    for name in ("output.out", "input.inp"):
        for f in done["tmp_path"].rglob(name):
            f.unlink()
    assert _analyse(session, done) == before


def test_the_photophysics_payload_has_an_esd_entry(tmp_path):
    from autodft.analysis.spectroscopy import analyze_spectra
    from autodft.db import get_session, init_db, reset_engine
    from autodft.models import ComputationHeader, Molecule, MoleculeState

    settings = Settings()
    settings.storage.data_path = str(tmp_path)
    reset_engine()
    try:
        init_db(settings)
        with get_session() as session:
            header = ComputationHeader(header_text="!B3LYP Opt Freq\n")
            mol = Molecule(smiles="O=CC=O", project_name="nho/p")
            session.add_all([header, mol])
            session.commit()
            ids = dict(optimization_header_id=header.id, singlepoint_header_id=header.id)
            session.add_all([
                MoleculeState(molecule_id=mol.id, description="S0", multiplicity=1, charge=0,
                              metadata_json=json.dumps({categories.ESD: True}), **ids),
                MoleculeState(molecule_id=mol.id, description="S1", multiplicity=1, charge=0,
                              metadata_json=json.dumps({"esd_role": "S1"}), **ids),
                MoleculeState(molecule_id=mol.id, description="T1", multiplicity=3, charge=0,
                              metadata_json=json.dumps({"esd_role": "T1"}), **ids),
            ])
            session.commit()
        [entry] = analyze_spectra("nho/p", use_cache=False)["molecules"]
        assert entry["esd"]["status"] == "waiting"
        assert "uvvis" not in entry and "ir" not in entry and "stage" not in entry
    finally:
        reset_engine()


# ----------------------------------------------------------------------
# I3: an FC channel whose SOCME prints as 0.00 flags the rate and the
# derived yields that rest on it.
# ----------------------------------------------------------------------

# t1s0.out with the ISC rate constant that a SOCME of 0.00 cm-1 gives.
ZERO_RATE_T1S0 = (FIXTURES / "t1s0.out").read_text().replace("1.598812e-01", "0.000000e+00")


def _zero_socme(session, esd):
    """Give the T1->S0 rate a job whose SOCME reads as 0.00 cm-1 (FC print precision)."""
    task = _rates(session, esd)[TaskType.esd_isc_t1s0]
    payload = json.loads(task.inputs_json)
    payload["computed"]["jobs"][0]["socme_cm"] = 0.0
    task.inputs_json = json.dumps(payload)
    session.add(task)
    job = session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one()
    Path(job.job_path, "output.out").write_text(_with_diagnostics(task, ZERO_RATE_T1S0))
    session.commit()
    return task


class TestZeroSocme:
    def test_fc_zero_socme_flags_the_channel_and_the_derived_group(self, session, done):
        _zero_socme(session, done)
        result = _analyse(session, done)
        assert result["rates"]["isc_t1_s0"]["rate_s"] == 0.0
        assert result["rates"]["isc_t1_s0"]["socme_zero"] is True
        assert any("SOCME below ORCA's print precision" in f for f in result["flags"])
        assert any("τ(T1)" in f and "T1→S0 ISC" in f for f in result["flags"])

    def test_ht_mode_does_not_flag_a_zero_socme_record(self, session, done):
        task = _zero_socme(session, done)
        payload = json.loads(task.inputs_json)
        payload["computed"]["jobs"][0]["ht"] = True
        task.inputs_json = json.dumps(payload)
        session.add(task)
        session.commit()
        result = _analyse(session, done)
        assert "socme_zero" not in result["rates"]["isc_t1_s0"]
        assert not any("SOCME below" in f for f in result["flags"])

    def test_nonzero_socme_is_not_flagged(self, session, done):
        result = _analyse(session, done)
        assert "socme_zero" not in result["rates"]["isc_t1_s0"]
        assert not any("SOCME below" in f for f in result["flags"])


# ----------------------------------------------------------------------
# A1: an HT sublevel row shows ORCA's own per-sublevel SOC.
# ----------------------------------------------------------------------

# Acetone T1->S0 (FC+HT): ORCA's "Reference SOC" of the Ms = -1, 0, +1 jobs.
ACETONE_SOC = ("-6.326655e-05, -5.262007e-07", "0.000000e+00, 2.303857e-04", "-6.326690e-05, 5.261551e-07")


def _ht_t1s0(session, esd, present=(1, 2, 3)):
    s1 = esd["states"]["S1"]
    s1.metadata_json = json.dumps({**json.loads(s1.metadata_json), categories.ESD_HT: True})
    task = _rates(session, esd)[TaskType.esd_isc_t1s0]
    payload = json.loads(task.inputs_json)
    # An older builder's record: the channel's SOCME repeated on every sublevel.
    payload["computed"] = {"combine": "mean", "jobs": [
        {"triplet": 1, "sublevel": ms, "dele_cm": 14444.5, "socme_cm": 31.0} for ms in (-1, 0, 1)
    ]}
    task.inputs_json = json.dumps(payload)
    session.add_all([s1, task])
    job = session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one()
    one = (FIXTURES / "t1s0.out").read_text()
    Path(job.job_path, "output.out").write_text("".join(
        f"$$$$$$$$$$$$$$$$  JOB NUMBER  {i} $$$$$$$$$$$$$$\n"
        + (one.replace("0.000000e+00, 5.261203e-08", soc) if i in present else "ORCA finished by error\n")
        for i, soc in enumerate(ACETONE_SOC, 1)
    ))
    session.commit()


def test_ht_sublevel_rows_carry_the_parsed_soc(session, done):
    _ht_t1s0(session, done)
    jobs = _analyse(session, done)["rates"]["isc_t1_s0"]["jobs"]
    assert [j["sublevel"] for j in jobs] == [-1, 0, 1]
    assert [j["socme_cm"] for j in jobs] == [
        pytest.approx(13.886, abs=1e-3), pytest.approx(50.564, abs=1e-3), pytest.approx(13.886, abs=1e-3),
    ]


def test_fc_rows_keep_the_builders_socme(session, done):
    [job] = _analyse(session, done)["rates"]["isc_t1_s0"]["jobs"]
    task = _rates(session, done)[TaskType.esd_isc_t1s0]
    assert job["socme_cm"] == json.loads(task.inputs_json)["computed"]["jobs"][0]["socme_cm"]


# ----------------------------------------------------------------------
# A2: a multi-job rate missing one job's rate is not combined.
# ----------------------------------------------------------------------

def test_a_missing_sublevel_leaves_the_rate_incomplete(session, done):
    _ht_t1s0(session, done, present=(1, 3))
    result = _analyse(session, done)
    entry = result["rates"]["isc_t1_s0"]
    assert entry["status"] == "incomplete" and entry["rate_s"] is None
    assert "no rate from job 2 triplet 1 sublevel 0" in entry["reason"]
    assert [j["sublevel"] for j in entry["jobs"]] == [-1, 1]
    assert any(f.startswith("isc_t1_s0: 2 rates for 3 ORCA jobs") for f in result["flags"])
    assert "tau_t1_us" not in result["derived"]
    assert result["status"] == "done"


def test_an_unnumbered_mismatch_is_incomplete_too(session, done):
    task = _rates(session, done)[TaskType.esd_isc]
    job = session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one()
    text = (FIXTURES / "isc_two_channels.out").read_text()
    end = "****ORCA ESD FINISHED WITHOUT ERROR****"
    Path(job.job_path, "output.out").write_text(text.split(end)[0] + end + "\n")
    entry = _analyse(session, done, detail=False)["rates"]["isc"]
    assert entry["status"] == "incomplete" and entry["rate_s"] is None
    assert entry["reason"].startswith("1 rates for 2 ORCA jobs;") and "no rate from" not in entry["reason"]
    assert "jobs" not in entry


# ----------------------------------------------------------------------
# FC diagnostics: flagged per channel, never part of the rate.
# ----------------------------------------------------------------------

def _rewrite_isc(session, esd, rates):
    task = _rates(session, esd)[TaskType.esd_isc]
    job = session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one()
    text = (FIXTURES / "isc_two_channels.out").read_text()
    Path(job.job_path, "output.out").write_text(_with_diagnostics(task, text, rates))
    session.commit()


def test_diagnostics_are_reported_but_not_combined(session, done):
    result = _analyse(session, done, detail=False)
    isc = result["rates"]["isc"]
    assert isc["rate_s"] == pytest.approx(9.521341e3)
    t1, t2 = isc["diagnostics"]
    assert t1["triplet"] == 1 and t1["k_inlinew_100"] == pytest.approx(9.521341e3)  # the FC main job
    assert (t1["k_inlinew_25"], t1["k_inlinew_250"], t1["k_tcutfreq_300"]) == (1e4, 1e4, 1e4)
    assert t2["k_inlinew_100"] == 0.0  # its negative main rate, clamped
    assert not t1["width_sensitive"] and not t1["soft_mode_dominated"]
    assert not any("width-sensitive" in f or "soft-mode" in f for f in result["flags"])
    assert (isc["ht"], isc["ht_reason"]) == (False, "Herzberg-Teller not requested")


def test_width_and_soft_mode_flags(session, done):
    _rewrite_isc(session, done, {(1, "inlinew_25"): 1e2, (1, "tcutfreq_300"): 5e2})
    result = _analyse(session, done, detail=False)
    [t1, _] = result["rates"]["isc"]["diagnostics"]
    assert t1["width_sensitive"] and t1["soft_mode_dominated"]
    assert any(f.startswith("isc T1: width-sensitive") for f in result["flags"])
    assert any(f.startswith("isc T1: soft-mode dominated") for f in result["flags"])
    assert result["rates"]["isc"]["rate_s"] == pytest.approx(9.521341e3)


def test_diagnostics_within_limits_are_not_flagged(session, done):
    # 3x apart exactly and a 10x drop exactly: neither is "more than".
    _rewrite_isc(session, done, {(1, "inlinew_25"): 1e4, (1, "inlinew_250"): 3e4,
                                 (1, "tcutfreq_300"): 9.521341e2})
    [t1, _] = _analyse(session, done, detail=False)["rates"]["isc"]["diagnostics"]
    assert not t1["width_sensitive"] and not t1["soft_mode_dominated"]


def test_negligible_channels_are_not_judged(session, done):
    # T2's rates stay below 1 s-1: a 100x spread there means nothing.
    _rewrite_isc(session, done, {(2, "inlinew_25"): 1e-3, (2, "inlinew_250"): 1e-1})
    [_, t2] = _analyse(session, done, detail=False)["rates"]["isc"]["diagnostics"]
    assert not t2["width_sensitive"]


def test_a_real_mixed_ht_and_fc_output(session, done):
    # Formaldehyde (ORCA 6.1.1): T1 by HT (three TDA sublevels), T2 by FC, then the FC diagnostics.
    task = _rates(session, done)[TaskType.esd_isc]
    payload = json.loads(task.inputs_json)
    payload["computed"] = json.loads((FIXTURES / "isc_tda_diagnostics.json").read_text())
    task.inputs_json = json.dumps(payload)
    session.add(task)
    job = session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one()
    Path(job.job_path, "output.out").write_text((FIXTURES / "isc_tda_diagnostics.out").read_text())
    session.commit()
    result = _analyse(session, done)
    isc = result["rates"]["isc"]
    assert isc["rate_s"] == pytest.approx(8.169736e4 + 5.582226 + 8.150821e4 + 1.198548e-02)
    assert [(j["triplet"], j.get("sublevel"), j["ht"]) for j in isc["jobs"]] == [
        (1, -1, True), (1, 0, True), (1, 1, True), (2, None, False),
    ]
    assert isc["jobs"][3]["socme_cm"] == pytest.approx(41.49)
    t1, t2 = isc["diagnostics"]
    assert t2["k_inlinew_100"] == pytest.approx(1.198548e-02) and t2["k_inlinew_25"] == 0.0
    assert not t2["width_sensitive"]  # below 1 s-1: not judged
    assert isc["ht_reason"] == "T1: SOCME 0.00 cm-1 < 5; T2: SOCME 41.49 cm-1 >= 5"
    assert not any("incomplete" in f or "ORCA jobs" in f for f in result["flags"])


def test_a_warning_repeated_by_sublevels_is_listed_once(session, done):
    task = _rates(session, done)[TaskType.esd_phosp]
    path = Path(session.exec(select(ComputationJob).where(ComputationJob.task_id == task.id)).one().job_path)
    text = (path / "output.out").read_text()
    (path / "output.out").write_text(re.sub(r"(The sum of K\*K is:\s+)\S+", r"\g<1>8.5", text))
    flags = _analyse(session, done)["flags"]
    assert sum("K*K" in f for f in flags) == 1
