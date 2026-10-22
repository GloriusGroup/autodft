"""The reactivity workbook: Export Project's hook, the export_reactivity job, and cleanup's kept .gbw."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from sqlmodel import select

from autodft import categories
from autodft.analysis.reactivity import reactivity_rows
from autodft.api import project_jobs
from autodft.db import get_session
from autodft.extraction import results
from autodft.extraction.extractor import PipelineExtractor
from autodft.models import (
    ComputationJob, ComputationTask, Molecule, MoleculeState, ProjectJobKind, TaskStatus, TaskType,
)
from autodft.models.result import JobResult
from tests.test_export_tree import PROJECT, _root, api, db  # noqa: F401 - fixtures

FIXTURES = Path(__file__).parent / "fixtures" / "orca"
RELAXED = (FIXTURES / "reactivity_relaxed_closed.out").read_text()
VERTICAL = (FIXTURES / "reactivity_vert_open.out").read_text()
ATOMS = 51
SP_KINDS = ["relaxed", "vert_ox", "vert_red", "vert_spin_change"]
VERTICALS = (TaskType.singlepoint_vert_ox, TaskType.singlepoint_vert_red, TaskType.singlepoint_vert_spin_change)


def _state(session, flagged=True) -> MoleculeState:
    mol = Molecule(smiles="Brc1ccccc1", project_name=PROJECT)
    session.add(mol)
    session.commit()
    metadata = {"imag_freq_tolerance": -10.0, **({categories.REACTIVITY: True} if flagged else {})}
    state = MoleculeState(molecule_id=mol.id, description="S0", multiplicity=1, charge=0,
                          metadata_json=json.dumps(metadata))
    session.add(state)
    session.commit()
    return state


def _task(session, settings, state, task_type, name, output: str, depends_on=None,
          record: Optional[dict] = None) -> ComputationTask:
    task = ComputationTask(task_type=task_type, status=TaskStatus.successful, state_id=state.id,
                           header_id=1, depends_on_task_id=depends_on, has_followups=False)
    session.add(task)
    session.commit()
    path = settings.comp_data_path / f"mol_{state.molecule_id}" / name / "job_1"
    path.mkdir(parents=True)
    (path / "output.out").write_text(output)
    for extra in ("input.inp", "input.gbw", "input.tmp"):
        (path / extra).write_text(extra)
    job = ComputationJob(task_id=task.id, attempt=1, job_path=str(path), success=True)
    session.add(job)
    session.commit()
    if record is not None:
        session.add(JobResult(job_id=job.id, task_id=task.id, job_path=job.job_path, finished_at=job.time_end,
                              parser_version=results.PARSER_VERSION, data_json=json.dumps(record)))
        session.commit()
    return task


def _conformer(session, settings, state, name, energy_shift=0.0, imaginary=()) -> ComputationTask:
    """An optimisation with its relaxed singlepoint and the three verticals, read from the fixtures."""
    opt = _task(session, settings, state, TaskType.optimization, f"opt_{name}", "opt",
                record={"energy": None, "free_energy_correction": 0.1 + energy_shift,
                        "imaginary": list(imaginary), "followed_root": None})
    _task(session, settings, state, TaskType.singlepoint, f"sp_{name}", RELAXED, depends_on=opt.id)
    for task_type in VERTICALS:
        _task(session, settings, state, task_type, f"{task_type.value}_{name}", VERTICAL, depends_on=opt.id)
    return opt


def _standard(settings, imaginary_first=(), first_lower=True):
    with get_session() as session:
        state = _state(session)
        first = _conformer(session, settings, state, "a", 0.0 if first_lower else 0.5, imaginary_first)
        second = _conformer(session, settings, state, "b", 0.5 if first_lower else 0.0)
        return state.molecule_id, first.id, second.id


def _workbook(settings):
    return load_workbook(_root(settings) / "p_reactivity.xlsx")


def _rows(ws) -> list[dict]:
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_export_writes_the_three_sheets(db):
    _standard(db)
    PipelineExtractor(PROJECT).export_tree(db.export_data_path, False, False)
    assert _workbook(db).sheetnames == ["molecular", "atomic", "wiberg"]
    assert (_root(db) / "p_reactivity.json").is_file()


def test_one_molecular_row_per_sp_kind(db):
    _standard(db)
    molecular = reactivity_rows(PROJECT)["molecular"]
    assert [r["sp_kind"] for r in molecular] == SP_KINDS
    assert [(r["charge"], r["multiplicity"]) for r in molecular] == [(0, 1), (1, 2), (-1, 2), (0, 3)]
    relaxed, ox = molecular[0], molecular[1]
    assert relaxed["energy_Eh"] == -6394.867848665785
    assert relaxed["polarizability_iso_au"] == 656.31703
    assert (relaxed["quad_xx_au"], relaxed["quad_xy_au"]) == (-136.108961855, 2.269079489)
    assert (relaxed["s2"], relaxed["stable"]) == (None, None)
    assert (ox["s2"], ox["stable"], ox["stability_lowest_Eh"]) == (0.787496, True, 0.01543417)


def test_atomic_rows_match_the_atom_count(db):
    _standard(db)
    rows = reactivity_rows(PROJECT)
    for kind in SP_KINDS:
        atoms = [r for r in rows["atomic"] if r["sp_kind"] == kind]
        assert [a["atom_1based"] for a in atoms] == list(range(1, ATOMS + 1))
    first = rows["atomic"][0]
    assert (first["element"], first["hirshfeld_charge"], first["npa_charge"]) == ("Br", -0.055841, 0.0577)
    ox = next(r for r in rows["atomic"] if r["sp_kind"] == "vert_ox")
    assert (ox["hirshfeld_spin"], ox["npa_spin"]) == (0.086022, 0.07955)
    wiberg = [r for r in rows["wiberg"] if r["sp_kind"] == "relaxed"]
    assert len(wiberg) == 74
    assert (wiberg[0]["i_1based"], wiberg[0]["j_1based"], wiberg[0]["wiberg_index"]) == (1, 2, 1.0417)


def test_only_the_reported_conformer_unless_all(db):
    _, first, second = _standard(db, first_lower=False)
    assert {r["opt_task_id"] for r in reactivity_rows(PROJECT)["molecular"]} == {second}
    assert {r["conformer"] for r in reactivity_rows(PROJECT)["atomic"]} == {2}
    every = reactivity_rows(PROJECT, all_conformers=True)
    assert {r["opt_task_id"] for r in every["molecular"]} == {first, second}


def test_excluded_conformer_is_absent(db):
    _, first, second = _standard(db, imaginary_first=[-15.0])
    for all_conformers in (False, True):
        rows = reactivity_rows(PROJECT, all_conformers=all_conformers)
        for sheet in ("molecular", "atomic", "wiberg"):
            assert {r["opt_task_id"] for r in rows[sheet]} == {second}


def test_export_tree_passes_the_all_conformers_choice(db):
    _, first, second = _standard(db)
    PipelineExtractor(PROJECT).export_tree(db.export_data_path, True, False)
    assert {r["opt_task_id"] for r in _rows(_workbook(db)["molecular"])} == {first, second}
    PipelineExtractor(PROJECT).export_tree(db.export_data_path, False, False)
    assert {r["opt_task_id"] for r in _rows(_workbook(db)["molecular"])} == {first}


def test_a_record_without_reactivity_blocks_says_so(db):
    with get_session() as session:
        state = _state(session)
        opt = _task(session, db, state, TaskType.optimization, "opt", "opt",
                    record={"energy": None, "free_energy_correction": 0.1, "imaginary": [], "followed_root": None})
        _task(session, db, state, TaskType.singlepoint, "sp", "gone", depends_on=opt.id,
              record={"energy": -1.0, "free_energy_correction": None, "npa_charges": None, "reactivity": None})
    rows = reactivity_rows(PROJECT)
    assert [(r["sp_kind"], r["energy_Eh"], r["note"]) for r in rows["molecular"]] == [
        ("relaxed", -1.0, "no reactivity output"),
    ]
    assert rows["atomic"] == rows["wiberg"] == []


def test_no_workbook_without_a_reactivity_molecule(db):
    with get_session() as session:
        state = _state(session, flagged=False)
        _conformer(session, db, state, "a")
    PipelineExtractor(PROJECT).export_tree(db.export_data_path, False, False)
    assert not (_root(db) / "p_reactivity.xlsx").exists()
    assert not (_root(db) / "p_reactivity.json").exists()


def test_export_reactivity_job(db):
    _standard(db)
    result = project_jobs._execute(ProjectJobKind.export_reactivity, PROJECT, {"all_conformers": True}, db)
    assert (result["format"], result["downloadable"]) == ("reactivity", True)
    assert Path(result["path"]) == _root(db) / "p_reactivity.xlsx"
    data = json.loads(Path(result["json_path"]).read_text())
    assert data == json.loads(json.dumps(reactivity_rows(PROJECT, all_conformers=True)))
    assert len(data["molecular"]) == 2 * len(SP_KINDS)


def test_export_reactivity_route(api, db):  # noqa: F811
    client, headers = api
    _standard(db)
    r = client.post("/api/projects/nho:p/export?format=reactivity&all_conformers=true", headers=headers)
    assert r.status_code == 202, r.text
    job = r.json()["job"]
    assert (job["kind"], job["params"]) == ("export_reactivity", {"all_conformers": True})
    project_jobs.join_all(timeout=10)
    assert (_root(db) / "p_reactivity.xlsx").is_file()


def test_export_reactivity_refused_when_archived(api, db):  # noqa: F811
    client, headers = api
    _standard(db)
    with get_session() as session:
        for mol in session.exec(select(Molecule)).all():
            mol.archived = True
            session.add(mol)
        session.commit()
    r = client.post("/api/projects/nho:p/export?format=reactivity", headers=headers)
    assert r.status_code == 409, r.text


def test_cleanup_keeps_gbw_of_reactivity_singlepoints(db):
    with get_session() as session:
        state = _state(session)
        _conformer(session, db, state, "a")
        _conformer(session, db, _state(session, flagged=False), "p")
        paths = {
            job.job_path: (task.task_type, task.state_id == state.id)
            for job, task in session.exec(
                select(ComputationJob, ComputationTask)
                .join(ComputationTask, ComputationJob.task_id == ComputationTask.id)
            ).all()
        }
    assert len(paths) == 10
    PipelineExtractor("__all__").cleanup_large_files()
    for path, (task_type, flagged) in paths.items():
        names = sorted(p.name for p in Path(path).iterdir())
        keeps_gbw = flagged and task_type.value in categories.REACTIVITY_TYPES
        assert ("input.gbw" in names) is keeps_gbw, (path, names)
        assert "input.tmp" not in names
